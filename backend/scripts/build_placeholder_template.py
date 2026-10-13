"""Builds a placeholder DO set template from do_set.yaml: one page per document,
each listing its values as {{ key }} placeholders. Replace it with the real set.

    python scripts/build_placeholder_template.py
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt

BACKEND_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_DIR))

from app.do_set import load_do_set  # noqa: E402


def main() -> None:
    do_set = load_do_set(BACKEND_DIR / "do_set.yaml")
    out = BACKEND_DIR / "templates" / "do_set.docx"
    document = Document()
    document.styles["Normal"].font.size = Pt(11)
    for number, page in enumerate(do_set.pages, start=1):
        document.add_heading(f"{number}. {page.title}", level=1)
        document.add_paragraph("Job {{ job_no }}        Date {{ today }}")
        document.add_paragraph("Placeholder layout. Replace this page with the real one and keep the placeholder names.")
        table = document.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for key in page.keys:
            cells = table.add_row().cells
            cells[0].text = do_set.fields[key].label
            cells[1].text = "{{ " + key + " }}"
        if number < len(do_set.pages):
            document.add_page_break()
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    print(f"Wrote {out.relative_to(BACKEND_DIR)} with {len(do_set.pages)} pages.")


if __name__ == "__main__":
    main()
